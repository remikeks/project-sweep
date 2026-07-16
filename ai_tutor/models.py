"""
Models for the SWEEP AI Tutor.

This is deliberately simple for the MVP: course/module text already
living on Course/CourseModule is used directly as context, and
CourseMaterial lets staff attach extra source documents (a lecture PDF,
slide deck, or transcript) whose text gets extracted once at upload time
and reused on every question. TutorInteraction is an audit log of what
was asked and answered, and doubles as the basis for a simple per-user
daily quota so a public-facing "Ask AI" button can't run up an unbounded
API bill.

Nothing here is agentic yet — there's no tool use, no autonomy, no
memory beyond this log. That's intentional: the brief for this iteration
is a course-aware Q&A/summary tutor, with room to grow into something
more capable later without a data-model rewrite.
"""

import io

from django.conf import settings
from django.db import models

from courses.models import Course, CourseModule

User = settings.AUTH_USER_MODEL

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".pptx", ".txt", ".md")


def course_material_upload_path(instance, filename):
    return f"course_materials/course_{instance.course_id}/{filename}"


class CourseMaterial(models.Model):
    """
    A source document (lecture slides, a transcript, a reading) the AI
    Tutor can draw on in addition to the course/module text already in
    the database. Attach it to a course only for course-wide context, or
    to a specific module to scope it to that lesson.
    """

    course = models.ForeignKey(Course, related_name="ai_materials", on_delete=models.CASCADE)
    module = models.ForeignKey(
        CourseModule,
        related_name="ai_materials",
        on_delete=models.CASCADE,
        null=True,
        blank=True,
        help_text="Leave blank to make this material available across the whole course.",
    )
    title = models.CharField(max_length=200)
    file = models.FileField(
        upload_to=course_material_upload_path,
        help_text="Supported for automatic text extraction: PDF, DOCX, PPTX, TXT, MD. "
        "If you replace the file later, use the 'Re-run text extraction' admin action.",
    )
    extracted_text = models.TextField(blank=True)
    char_count = models.PositiveIntegerField(default=0)
    extraction_error = models.CharField(max_length=500, blank=True)
    is_active = models.BooleanField(default=True)
    uploaded_by = models.ForeignKey(
        User, null=True, blank=True, on_delete=models.SET_NULL, related_name="uploaded_materials"
    )
    uploaded_at = models.DateTimeField(auto_now_add=True)

    class Meta:
        ordering = ["-uploaded_at"]

    def __str__(self):
        return self.title

    def save(self, *args, **kwargs):
        is_new_file = self.pk is None or not self.extracted_text
        super().save(*args, **kwargs)
        if is_new_file and self.file:
            self.extract_text()
            super().save(update_fields=["extracted_text", "char_count", "extraction_error"])

    def extract_text(self):
        """Populate extracted_text from the uploaded file. Safe to re-run."""
        name = self.file.name.lower()
        text = ""
        error = ""
        try:
            with self.file.open("rb") as fh:
                data = fh.read()

            if name.endswith(".pdf"):
                text = _extract_pdf(data)
            elif name.endswith(".docx"):
                text = _extract_docx(data)
            elif name.endswith(".pptx"):
                text = _extract_pptx(data)
            elif name.endswith((".txt", ".md")):
                text = data.decode("utf-8", errors="ignore")
            else:
                error = (
                    f"Unsupported file type for automatic extraction. "
                    f"Supported: {', '.join(SUPPORTED_EXTENSIONS)}"
                )
        except Exception as exc:  # noqa: BLE001 - surface any extraction failure to the admin
            error = f"Extraction failed: {exc}"

        self.extracted_text = text.strip()
        self.char_count = len(self.extracted_text)
        self.extraction_error = error


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(data: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs if p.text)


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    chunks = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text:
                chunks.append(shape.text_frame.text)
    return "\n".join(chunks)


class TutorInteraction(models.Model):
    """
    A single AI Tutor exchange: either a free-form question and its
    answer, or a one-click summary. Kept for basic auditing/analytics
    and to enforce the daily quota in ai_tutor.services.
    """

    class Kind(models.TextChoices):
        CHAT = "chat", "Question"
        SUMMARY = "summary", "Summary"

    user = models.ForeignKey(User, related_name="tutor_interactions", on_delete=models.CASCADE)
    course = models.ForeignKey(Course, related_name="tutor_interactions", on_delete=models.CASCADE)
    module = models.ForeignKey(
        CourseModule, related_name="tutor_interactions", on_delete=models.CASCADE, null=True, blank=True
    )
    kind = models.CharField(max_length=10, choices=Kind.choices, default=Kind.CHAT)
    question = models.TextField(blank=True)
    answer = models.TextField(blank=True)
    was_error = models.BooleanField(default=False)
    created_at = models.DateTimeField(auto_now_add=True)

    class Meta:
        ordering = ["-created_at"]

    def __str__(self):
        scope = self.module.title if self.module else self.course.title
        return f"{self.get_kind_display()} · {self.user} · {scope}"
