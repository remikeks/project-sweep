"""Build a new, non-destructive research-informed expansion candidate.

One structured source drives DOCX, PDF and deck inputs. This command never
uploads, approves, publishes, activates courses or creates assessment IDs.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from courses.curriculum_validation import validate_catalogue

SOURCE = ROOT / "production_courseware/expansion_review/teaching_source.json"
ASSESSMENTS = ROOT / "production_courseware/expansion_review/assessment_source.json"
NOTICE = ("Use the laws, professional scope, organisational policies and approved referral "
          "routes that apply to your setting. This course is educational and does not replace "
          "urgent safeguarding, clinical or legal procedures. All worked cases are fictional. "
          "Research notes distinguish published findings from SWEEP teaching examples.")


def load_sources():
    source = json.loads(SOURCE.read_text(encoding="utf-8"))
    banks = json.loads(ASSESSMENTS.read_text(encoding="utf-8"))
    if [c["code"] for c in source["courses"]] != [c["code"] for c in banks["courses"]]:
        raise ValueError("Course source and assessment identities differ")
    references = {r["id"]: r for r in source["references"]}
    if len(references) != len(source["references"]):
        raise ValueError("Duplicate reference IDs")
    catalogue = []
    for course, bank in zip(source["courses"], banks["courses"]):
        entries = json.loads((ROOT / "production_courseware/drafts" / bank["draft"] / "course_catalogue.json").read_text(encoding="utf-8"))
        if len(entries) != 1 or entries[0]["code"] != course["code"]:
            raise ValueError("Unexpected course identity")
        catalogue.extend(entries)
        if [m["order"] for m in course["modules"]] != [1, 2, 3]:
            raise ValueError("Expected exactly three ordered modules")
        if not all(key in references for key in course["references"]):
            raise ValueError("Unresolved reference")
        for module, assessed in zip(course["modules"], bank["modules"]):
            if assessed["order"] != module["order"] or len(assessed["items"]) != 5:
                raise ValueError("Assessment coverage mismatch")
            for field in ("heading", "example", "activity", "feedback", "slide"):
                if not module[field].strip():
                    raise ValueError(f"Empty {field}")
            if len(module["explanation"]) < 2:
                raise ValueError("Missing teaching explanation")
    validate_catalogue(catalogue, expected_school_count=2)
    return source, banks, catalogue


def clean_heading(value):
    return " ".join(re.sub(r"[^\w\s]", " ", value.replace("&", "and")).split())


def reference_blocks(course, references, *, evidence_notes=True):
    blocks = [("h1", "Sources and evidence limits")]
    for key in course["references"]:
        ref = references[key]
        blocks.append(("p", f"[{key}] {ref['citation']} {ref['url']}"))
        if evidence_notes:
            blocks.append(("p", f"Use in this course: {ref['summary']} Limit: {ref['limit']}"))
    blocks.append(("p", "The index record is a discovery aid, not a guarantee of quality. Read the source before applying its findings to a different setting. No article endorses this course."))
    return blocks


def guide_blocks(course, bank, catalogue, references):
    blocks = [("title", clean_heading(catalogue["title"]) + " learner guide"),
              ("p", f"{course['code']} | Version 1.1.0 | SWEEP Academy"),
              ("p", catalogue["outcome"]), ("p", NOTICE),
              ("h1", "How to use this guide"),
              ("p", "Read each lesson, explain the worked example in your own words and complete its activity before checking the feedback. The decision points cover the concepts assessed; they are not a substitute for the explanations. Use the separate worksheet to integrate the three lessons. Study time depends on discussion and local adaptation, not a guaranteed completion duration."),
              ("h1", "Practice case"), ("p", course["case"])]
    for module, assessed, meta in zip(course["modules"], bank["modules"], catalogue["modules"]):
        blocks += [("break", ""), ("h1", clean_heading(f"Module {module['order']} {meta['title']}")),
                   ("p", assessed["outcome"]), ("h2", module["heading"])]
        blocks.extend(("p", paragraph) for paragraph in module["explanation"])
        blocks += [("h2", "Worked example"), ("p", module["example"]),
                   ("h2", "Decision points")]
        blocks.extend(("p", item[0]) for item in assessed["items"])
        blocks += [("h2", "Apply the lesson"), ("p", module["activity"]),
                   ("h2", "Check your reasoning"), ("p", module["feedback"])]
    blocks += [("break", ""), ("h1", "Assessment preparation"),
               ("p", "Explain the reasons behind each decision point. Practise with the worksheet using a different situation from the worked examples. The proposed assessment has 15 single-answer questions and a 70% pass threshold, requiring at least 11 correct whole-mark answers. Published settings and availability must be verified in the learning platform; this guide does not establish an assessment ID or certify completion.")]
    blocks += reference_blocks(course, references)
    return blocks


def worksheet_blocks(course, bank, catalogue, references):
    blocks = [("title", clean_heading(catalogue["title"]) + " practice worksheet"),
              ("p", f"{course['code']} | Version 1.1.0 | SWEEP Academy"),
              ("p", "Apply all three lessons to the fictional case below. Do not use identifiable service-user information. Write a reasoned plan that makes uncertainty and decision authority explicit."),
              ("p", NOTICE), ("h1", "Practice case"), ("p", course["case"])]
    for module in course["modules"]:
        if module["order"] == 2:
            blocks += [("break", "")]
        blocks += [("h2", f"Activity {module['order']}"), ("p", module["activity"]), ("space", "120")]
    blocks += [("break", ""), ("h1", "Self assessment and facilitator discussion"),
               ("p", "For each activity, use 0 for missing reasoning, 1 for a partly justified response and 2 for a justified response with relevant safeguards and limits. This is a formative rubric, not the platform's certification score.")]
    for module in course["modules"]:
        blocks += [("h2", f"Activity {module['order']} criteria"), ("p", module["feedback"])]
    blocks += [("h2", "Next learning step"), ("p", "Identify one assumption you changed and one procedure you must verify locally. Explain who can resolve it before you apply the plan in practice.")]
    blocks += reference_blocks(course, references, evidence_notes=False)
    return blocks


def transcript_blocks(course, bank, catalogue, module, assessed, references):
    blocks = [("title", clean_heading(module["heading"]) + " teaching script"),
              ("p", f"{course['code']} | Module {module['order']} | Version 1.1.0"),
              ("p", "This is a written lesson plan and narration script, not a recorded video. The facilitator should rehearse timing, provide accessible captions if recording, and verify local procedures before delivery."),
              ("h1", "Learning purpose"), ("p", assessed["outcome"]), ("p", NOTICE),
              ("h1", "Opening narration"), ("p", f"This lesson examines how to {assessed['outcome'][0].lower() + assessed['outcome'][1:]}")]
    blocks += [("h1", "Core narration")]
    blocks.extend(("p", paragraph) for paragraph in module["explanation"])
    blocks += [("h1", "Worked example narration"), ("p", module["example"]),
               ("h1", "Pause for learner activity"), ("p", module["activity"]),
               ("p", "Pause for a written response or discussion. Do not present private case information on screen."),
               ("h1", "Debrief narration"), ("p", module["feedback"]),
               ("h1", "Decision point recap")]
    blocks.extend(("p", item[0]) for item in assessed["items"])
    blocks += [("h1", "Recording plan"),
               ("p", "Show the lesson title, a short decision sequence and the fictional example. Read key visual information aloud. Separate the activity from the debrief and make pauses editable. Check names, acronyms and captions manually; do not claim a video exists until a recording has been produced.")]
    blocks += reference_blocks(course, references, evidence_notes=False)
    return blocks


def write_docx(blocks, destination):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.7)
    sec.left_margin = sec.right_margin = Inches(.8)
    for name, size in (("Normal", 11), ("Title", 24), ("Heading 1", 17), ("Heading 2", 13)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
        for border in style.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.0
        if name != "Normal":
            style.paragraph_format.keep_with_next = True
    for kind, text in blocks:
        if kind == "break":
            doc.add_page_break()
        elif kind == "space":
            p = doc.add_paragraph("Write your response here")
            p.paragraph_format.space_after = Pt(int(text))
        else:
            doc.add_paragraph(text, style={"title": "Title", "h1": "Heading 1", "h2": "Heading 2", "p": "Normal"}[kind])
    doc.core_properties.author = "SWEEP Academy"
    doc.core_properties.title = blocks[0][1]
    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)


def write_pdf(blocks, destination):
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="BodyTextSWEEP", fontName="Helvetica", fontSize=10.5, leading=14, spaceAfter=6))
    styles.add(ParagraphStyle(name="TitleSWEEP", fontName="Helvetica-Bold", fontSize=23, leading=28, spaceAfter=15, textColor=colors.black))
    for name in ("Heading1", "Heading2"):
        styles[name].textColor = colors.black
        styles[name].keepWithNext = True
        styles[name].spaceBefore = 9
        styles[name].spaceAfter = 6
    story = []
    for kind, text in blocks:
        if kind == "break":
            story.append(PageBreak())
        elif kind == "space":
            story += [Paragraph("Write your response here", styles["BodyTextSWEEP"]), Spacer(1, int(text))]
        else:
            story.append(Paragraph(escape(text), styles[{"title":"TitleSWEEP", "h1":"Heading1", "h2":"Heading2", "p":"BodyTextSWEEP"}[kind]]))
    destination.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(destination), pagesize=letter, rightMargin=57.6, leftMargin=57.6,
                      topMargin=50.4, bottomMargin=50.4, title=blocks[0][1], author="SWEEP Academy").build(story)


def build(output, *, revise_unpublished=False):
    source, banks, catalogue = load_sources()
    if output.exists():
        receipt = output / "evidence_receipt.json"
        if not revise_unpublished or not receipt.is_file() or json.loads(receipt.read_text())["published"] is not False:
            raise FileExistsError("Existing output requires an explicitly selected unpublished generated candidate")
    output.mkdir(parents=True, exist_ok=True)
    references = {r["id"]: r for r in source["references"]}
    coverage = []
    for course, bank, meta in zip(source["courses"], banks["courses"], catalogue):
        prefix = f"{meta['number']:02d}-{meta['slug']}"
        for name, blocks in (("learner-guide", guide_blocks(course, bank, meta, references)),
                             ("practice-worksheet", worksheet_blocks(course, bank, meta, references))):
            write_docx(blocks, output / "docx" / f"{prefix}-{name}.docx")
            write_pdf(blocks, output / "pdf" / f"{prefix}-{name}.pdf")
        for module, assessed in zip(course["modules"], bank["modules"]):
            write_docx(transcript_blocks(course, bank, meta, module, assessed, references),
                       output / "docx" / f"{prefix}-module-{module['order']}-video-plan-transcript.docx")
            coverage.append({"course":course["code"], "module":module["order"], "taught_decision_points":len(assessed["items"]), "activity":True, "worked_example":True})
    (output / "course_catalogue.json").write_text(json.dumps(catalogue, indent=2), encoding="utf-8")
    (output / "evidence_receipt.json").write_text(json.dumps({
        "version":source["version"], "teaching_source_sha256":hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
        "assessment_source_sha256":hashlib.sha256(ASSESSMENTS.read_bytes()).hexdigest(),
        "owner_authorized":True, "independent_subject_matter_review":False,
        "visual_review_complete":False, "published":False, "coverage":coverage,
    }, indent=2), encoding="utf-8")
    print(f"Built 20 DOCX and 8 PDF files at {output}. Deck generation and visual checks remain required.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--revise-unpublished", action="store_true")
    args = parser.parse_args()
    build(args.output.resolve(), revise_unpublished=args.revise_unpublished)
