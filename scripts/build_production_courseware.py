"""Build the SWEEP production courseware package from the catalogue outline.

The command intentionally produces general, internationally usable learning
materials. It does not embed country-specific reporting routes, clinical
protocols, legal advice, or emergency numbers. Those details belong in the
approved local release for each delivery jurisdiction.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
from dataclasses import asdict, dataclass
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


RELEASE_VERSION = "1.0.0"
NAVY = RGBColor(11, 44, 122)
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(89, 89, 89)
COURSE_CODE_PREFIX = "SWP"


@dataclass
class Module:
    order: int
    title: str
    focus: str


@dataclass
class Course:
    number: int
    school: str
    title: str
    outcome: str
    modules: list[Module]
    assets: list[str]
    cbt_focus: list[str]

    @property
    def code(self) -> str:
        return f"{COURSE_CODE_PREFIX}-{self.number:03d}"

    @property
    def slug(self) -> str:
        return slugify(self.title)


def slugify(value: str) -> str:
    value = value.lower().replace("&", "and")
    value = re.sub(r"[^a-z0-9]+", "-", value).strip("-")
    return value


def clean_markdown(value: str) -> str:
    """The outline uses bold markers for labels; learner assets use plain text."""
    return value.replace("**", "").strip()


def parse_outlines(source: Path) -> list[Course]:
    """Read the stable heading structure in COURSE_CONTENT_OUTLINES.md."""
    lines = source.read_text(encoding="utf-8").splitlines()
    school = ""
    courses: list[Course] = []
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if line.startswith("## School of "):
            school = line[3:]
            index += 1
            continue
        match = re.fullmatch(r"### (\d+)\. (.+)", line)
        if not match:
            index += 1
            continue
        number = int(match.group(1))
        title = match.group(2)
        block: list[str] = []
        index += 1
        while index < len(lines) and not lines[index].startswith("### ") and not lines[index].startswith("## "):
            block.append(lines[index].strip())
            index += 1
        text = "\n".join(block)
        outcome_match = re.search(r"\*\*Outcome:\*\*\s*(.+)", text)
        asset_match = re.search(r"\*\*Assets:\*\*\s*(.+?)\s+\*\*CBT focus:\*\*\s*(.+)", text)
        modules: list[Module] = []
        for item in block:
            module_match = re.fullmatch(r"(\d+)\.\s+(.+?)\s+[—-]\s+(.+)", item)
            if module_match:
                modules.append(
                    Module(
                        order=int(module_match.group(1)),
                        title=clean_markdown(module_match.group(2)),
                        focus=module_match.group(3).rstrip(".").strip(),
                    )
                )
        if not school or not outcome_match or len(modules) != 3 or not asset_match:
            raise ValueError(f"Could not parse a complete course brief for {title!r}.")
        courses.append(
            Course(
                number=number,
                school=school,
                title=title,
                outcome=outcome_match.group(1).strip(),
                modules=modules,
                assets=[item.strip() for item in asset_match.group(1).split(";") if item.strip()],
                cbt_focus=[item.strip() for item in asset_match.group(2).rstrip(".").split(",") if item.strip()],
            )
        )
    if len(courses) != 30:
        raise ValueError(f"Expected 30 courses in the outline, found {len(courses)}.")
    return courses


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = tc_pr.find(qn("w:shd"))
    if shade is None:
        shade = OxmlElement("w:shd")
        tc_pr.append(shade)
    shade.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_run_font(run, *, size: float | None = None, bold: bool | None = None, color=None) -> None:
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = color


def setup_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    for style_name, size in (("Normal", 10.8), ("Title", 25), ("Heading 1", 17), ("Heading 2", 13)):
        style = doc.styles[style_name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style.font.size = Pt(size)
        style.font.color.rgb = BLACK
        style.font.bold = style_name != "Normal"
        style.paragraph_format.space_after = Pt(7 if style_name == "Normal" else 10)
        style.paragraph_format.keep_with_next = style_name != "Normal"
    if "Small Note" not in [style.name for style in doc.styles]:
        note = doc.styles.add_style("Small Note", WD_STYLE_TYPE.PARAGRAPH)
        note.font.name = "Aptos"
        note.font.size = Pt(9)
        note.font.color.rgb = GRAY
        note.paragraph_format.space_after = Pt(6)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("SWEEP Academy learning materials")
    set_run_font(run, size=8.5, color=GRAY)


def add_title(doc: Document, title: str, subtitle: str) -> None:
    paragraph = doc.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(title)
    set_run_font(run, size=25, bold=True, color=BLACK)
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(16)
    run = sub.add_run(subtitle)
    set_run_font(run, size=12, color=GRAY)


def add_bullets(doc: Document, entries: list[str]) -> None:
    for entry in entries:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(entry)
        set_run_font(run, size=10.8, color=BLACK)


def focus_list(module: Module) -> list[str]:
    value = module.focus.replace(" and ", ", ")
    parts = [part.strip() for part in re.split(r",\s*", value) if part.strip()]
    return parts


def localisation_note(doc: Document) -> None:
    paragraph = doc.add_paragraph(style="Small Note")
    run = paragraph.add_run(
        "Practice notice: Apply the law, organisational policy, supervision requirements, and approved referral pathways that govern your setting. "
        "This material supports learning and does not replace urgent safeguarding, clinical, or legal procedures."
    )
    set_run_font(run, size=9, color=GRAY)


def create_learner_guide(course: Course, destination: Path) -> None:
    doc = Document()
    setup_document(doc)
    add_title(doc, course.title, f"Learner guide  |  {course.code}  |  {course.school}")
    opening = doc.add_paragraph()
    opening.paragraph_format.space_after = Pt(10)
    run = opening.add_run(
        f"This guide supports a three-part learning sequence for {course.title}. "
        f"By the end of the course, you should be able to {course.outcome[0].lower() + course.outcome[1:]}"
    )
    set_run_font(run, size=11.5, color=BLACK)
    localisation_note(doc)
    doc.add_heading("Learning objectives", level=1)
    add_bullets(doc, [course.outcome, *[f"Apply the key ideas from {module.title}." for module in course.modules]])
    doc.add_heading("Course structure", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ["Module", "Focus", "Suggested study time"]
    for cell, value in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, "0B2C7A")
        set_cell_borders(cell)
        cell.vertical_alignment = 1
        run = cell.paragraphs[0].add_run(value)
        set_run_font(run, size=9.5, bold=True, color=RGBColor(255, 255, 255))
    for module in course.modules:
        cells = table.add_row().cells
        values = [f"{module.order}. {module.title}", module.focus, "10 to 15 minutes"]
        for index, (cell, value) in enumerate(zip(cells, values)):
            set_cell_borders(cell)
            cell.vertical_alignment = 1
            if module.order % 2 == 0:
                set_cell_shading(cell, "F2F6FA")
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 2 else WD_ALIGN_PARAGRAPH.LEFT
            run = cell.paragraphs[0].add_run(value)
            set_run_font(run, size=9.2, color=BLACK)
    for module in course.modules:
        doc.add_page_break()
        doc.add_heading(f"Module {module.order}  {module.title}", level=1)
        intro = doc.add_paragraph()
        run = intro.add_run(
            f"This module focuses on {module.focus}. Work through the material with a current practice setting in mind, "
            "but remove identifying information from any examples you record."
        )
        set_run_font(run, size=11, color=BLACK)
        doc.add_heading("What to pay attention to", level=2)
        add_bullets(doc, [item[:1].upper() + item[1:] for item in focus_list(module)])
        doc.add_heading("Learning sequence", level=2)
        paragraphs = [
            f"Start by defining the purpose of {module.title.lower()} in relation to the course outcome. Notice where professional judgement is needed and where a local procedure must guide the next step.",
            f"Examine the practical implications of {module.focus}. Separate observations from assumptions, identify strengths as well as concerns, and be clear about information that must be checked with a person, family, colleague, or approved source.",
            "Use a respectful, collaborative approach. Explain your role, seek informed participation where appropriate, and document only information that is relevant, accurate, and necessary for the purpose of the work.",
            "When a concern exceeds your role, supervision, competence, or authority, pause and use the correct escalation route. A prompt referral and a clear handover can protect continuity of support.",
        ]
        for text in paragraphs:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(8)
            run = paragraph.add_run(text)
            set_run_font(run, size=10.8, color=BLACK)
        doc.add_heading("Practice reflection", level=2)
        prompt = doc.add_paragraph()
        run = prompt.add_run(
            f"Choose a de-identified practice situation. What information would help you work through {module.title.lower()} safely, and what would require consultation or a local policy check?"
        )
        set_run_font(run, size=10.8, color=BLACK)
    doc.add_page_break()
    doc.add_heading("Course activity", level=1)
    activity = doc.add_paragraph()
    run = activity.add_run(
        "Use the accompanying worksheet to apply the course sequence to a short case. Your completed activity should show what you noticed, the action you would take within your role, and what you would escalate or verify locally."
    )
    set_run_font(run, size=10.8, color=BLACK)
    doc.add_heading("Assessment preparation", level=1)
    add_bullets(doc, [f"Review the distinction between: {focus}." for focus in course.cbt_focus])
    doc.add_heading("References and local adaptation", level=1)
    references = [
        "SWEEP Academy Course Content Outlines, production blueprint for this course.",
        "The current legislation, professional standards, safeguarding procedures, and referral directories for the delivery jurisdiction.",
        "The learner's employer or placement policy, supervision arrangements, and documentation standards.",
    ]
    add_bullets(doc, references)
    doc.core_properties.title = f"{course.title} Learner Guide"
    doc.core_properties.author = "SWEEP Academy"
    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)


def create_worksheet(course: Course, destination: Path) -> None:
    doc = Document()
    setup_document(doc)
    add_title(doc, course.title, f"Practice worksheet  |  {course.code}")
    intro = doc.add_paragraph()
    run = intro.add_run(
        "Use this de-identified learning activity after completing the three modules. Record only what is necessary for your learning. Do not include client, family, patient, student, or colleague names."
    )
    set_run_font(run, size=11, color=BLACK)
    localisation_note(doc)
    doc.add_heading("Practice scenario", level=1)
    scenario = doc.add_paragraph()
    run = scenario.add_run(
        f"You are supporting a person, family, group, or community that needs help related to {course.title.lower()}. "
        "The situation contains incomplete information, competing priorities, and at least one issue that may need advice from a supervisor or partner service."
    )
    set_run_font(run, size=10.8, color=BLACK)
    doc.add_heading("Activity instructions", level=1)
    add_bullets(doc, [
        "Work through the modules in order and use the questions below to organise your response.",
        "Distinguish information you know from information you need to verify.",
        "Identify the immediate action that sits within your role and any decision that needs escalation.",
        "Use your local policy and referral information before applying this activity in practice.",
    ])
    prompts = [
        ("1. What matters now", f"What information, strengths, risks, and priorities would you consider when approaching {course.modules[0].title.lower()}?"),
        ("2. Safe and collaborative response", f"How would you apply {course.modules[1].title.lower()} while communicating respectfully and within your scope?"),
        ("3. Coordination and follow-up", f"What documentation, review point, consultation, referral, or handover would be appropriate when applying {course.modules[2].title.lower()}?"),
        ("4. Reflection", "What assumption did you need to test, and how would you check it without causing harm or excluding the person's voice?"),
    ]
    for heading, question in prompts:
        doc.add_heading(heading, level=2)
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(question)
        set_run_font(run, size=10.5, color=BLACK)
        response_table = doc.add_table(rows=1, cols=1)
        response_table.style = "Table Grid"
        cell = response_table.cell(0, 0)
        set_cell_borders(cell)
        set_cell_shading(cell, "FAFBFC")
        cell.height = Inches(1.05)
        cell.vertical_alignment = 0
        marker = cell.paragraphs[0].add_run("Response")
        set_run_font(marker, size=9, color=GRAY)
    doc.add_heading("Self-check before submission", level=1)
    add_bullets(doc, [
        "My response uses de-identified information.",
        "I have separated observations from assumptions.",
        "I have named the action within my role and the action that needs consultation or escalation.",
        "I have considered equity, dignity, participation, and confidentiality.",
    ])
    doc.core_properties.title = f"{course.title} Practice Worksheet"
    doc.core_properties.author = "SWEEP Academy"
    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)


def create_transcript(course: Course, module: Module, destination: Path) -> None:
    doc = Document()
    setup_document(doc)
    add_title(doc, module.title, f"Video plan and transcript  |  {course.title}  |  Module {module.order}")
    localisation_note(doc)
    doc.add_heading("Lesson plan", level=1)
    plan = [
        ("Audience", "Learners completing the SWEEP Academy course sequence"),
        ("Duration", "Approximately 10 minutes"),
        ("Learning purpose", f"Help learners apply {module.title.lower()} to the course outcome."),
        ("Visual approach", "Use a simple title card, a neutral de-identified practice scenario, and accessible captions."),
        ("Facilitator preparation", "Check the current local policy, escalation routes, and support contacts before recording or delivery."),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in plan:
        cells = table.add_row().cells
        for cell in cells:
            set_cell_borders(cell)
            cell.vertical_alignment = 1
        set_cell_shading(cells[0], "EAF1F8")
        label_run = cells[0].paragraphs[0].add_run(label)
        set_run_font(label_run, size=9.5, bold=True, color=BLACK)
        value_run = cells[1].paragraphs[0].add_run(value)
        set_run_font(value_run, size=9.5, color=BLACK)
    doc.add_heading("Caption-ready transcript", level=1)
    sections = [
        ("Opening 0:00 to 1:00", f"Welcome to module {module.order}, {module.title}. In this lesson, we focus on {module.focus}. The aim is to help you connect the topic to the course outcome: {course.outcome}"),
        ("Set the frame 1:00 to 3:00", f"Begin with the purpose of this work. {module.title} calls for careful attention to context, the person's perspective, available strengths, and the limits of your role. The details will vary by setting, so use local guidance whenever a decision has legal, safeguarding, clinical, or urgent consequences."),
        ("Core ideas 3:00 to 6:00", f"The key areas in this module are {module.focus}. Treat these as connected questions rather than a checklist. Gather relevant information, explain what you are doing, and record only what is necessary. If information is unclear, say what you know, what you need to confirm, and who can help you confirm it."),
        ("Practice example 6:00 to 8:30", f"Imagine a de-identified situation in which a person or community needs support related to {course.title.lower()}. Start by listening and checking immediate priorities. Consider what action sits within your role. Then identify what requires consultation, a formal referral, a mandated process, or a handover to another service."),
        ("Close and reflection 8:30 to 10:00", f"Before the next module, pause and reflect: what would you need to know to approach {module.title.lower()} safely and respectfully? Use the worksheet to write a short response. The next step is to connect this learning with {course.modules[module.order % len(course.modules)].title.lower()}.")
    ]
    for heading, body in sections:
        doc.add_heading(heading, level=2)
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(body)
        set_run_font(run, size=10.8, color=BLACK)
    doc.core_properties.title = f"{course.title} Module {module.order} Video Plan and Transcript"
    doc.core_properties.author = "SWEEP Academy"
    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)


def add_asset(
    entries: list[dict],
    course: Course,
    path: Path,
    asset_type: str,
    title: str,
    module_order: int | None = None,
) -> None:
    entries.append(
        {
            "course_number": course.number,
            "course_code": course.code,
            "course_slug": course.slug,
            "school": course.school,
            "module_order": module_order,
            "title": title,
            "asset_type": asset_type,
            "path": path.as_posix(),
            "version": RELEASE_VERSION,
            "language": "en",
            "publication_status": "in_review",
        }
    )


def build_release(source: Path, output: Path) -> None:
    courses = parse_outlines(source)
    docx_dir = output / "docx"
    if docx_dir.exists():
        shutil.rmtree(docx_dir)
    manifest: list[dict] = []
    content_catalogue: list[dict] = []
    for course in courses:
        prefix = f"{course.number:02d}-{course.slug}"
        guide = docx_dir / f"{prefix}-learner-guide.docx"
        worksheet = docx_dir / f"{prefix}-practice-worksheet.docx"
        create_learner_guide(course, guide)
        create_worksheet(course, worksheet)
        add_asset(manifest, course, guide.relative_to(output), "learner_guide", f"{course.title} learner guide DOCX")
        add_asset(manifest, course, worksheet.relative_to(output), "worksheet", f"{course.title} practice worksheet DOCX")
        for module in course.modules:
            transcript = docx_dir / f"{prefix}-module-{module.order}-video-plan-transcript.docx"
            create_transcript(course, module, transcript)
            add_asset(
                manifest,
                course,
                transcript.relative_to(output),
                "transcript",
                f"{course.title} module {module.order} video plan and transcript DOCX",
                module.order,
            )
        catalogue_entry = asdict(course)
        catalogue_entry["code"] = course.code
        catalogue_entry["slug"] = course.slug
        content_catalogue.append(catalogue_entry)
    output.mkdir(parents=True, exist_ok=True)
    (output / "course_catalogue.json").write_text(json.dumps(content_catalogue, indent=2), encoding="utf-8")
    (output / "asset_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    with (output / "asset_manifest.csv").open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(manifest[0]))
        writer.writeheader()
        writer.writerows(manifest)
    print(f"Built {len(courses)} courses and {len(manifest)} DOCX assets in {output}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build_release(args.source.resolve(), args.output.resolve())


if __name__ == "__main__":
    main()
